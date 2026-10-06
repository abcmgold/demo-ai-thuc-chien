"""Xuất lịch trình hiện tại thành workbook Excel trong bộ nhớ."""

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CURRENCY_FORMAT = '#,##0" VNĐ"'


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]],
           widths: list[int], currency_columns: tuple[int, ...] = ()) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            # Nội dung mô tả được giữ là văn bản kể cả khi bắt đầu bằng dấu '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="173C32")
                cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
            else:
                cell.font = Font(name="Calibri", color="203128")
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F5F3EC")
                if cell.column in currency_columns and isinstance(cell.value, (int, float)):
                    cell.number_format = CURRENCY_FORMAT
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 32


def export_plan_excel(plan: dict[str, Any]) -> bytes:
    """Giữ số tiền dạng số, xuất đầy đủ mọi ngày và cảnh báo của kế hoạch."""
    request = plan["request"]
    total, budget = plan["total_vnd"], request["group_budget_vnd"]
    over_budget = max(0, total - budget)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Lịch trình du lịch Đắk Lắk"
    workbook.properties.creator = "Trợ lý Du lịch Đắk Lắk AI"
    summary_rows = [
        ["Hành trình", plan["summary"]],
        ["Số người", request["people"]],
        ["Số ngày", request["days"]],
        ["Số đêm", request["nights"]],
        ["Sở thích", " · ".join(request["interests"])],
        ["Điểm bắt đầu", request["start_area"]],
        ["Phương tiện mong muốn", request["transport"]],
        ["Mức tiện nghi", request["comfort"]],
        ["Nhịp độ", request["pace"]],
        ["Số bữa/ngày", request["meals_per_day"]],
        ["Ngân sách tính cho", request["budget_basis"]],
        ["Ngân sách nhập (VNĐ)", request["budget"]],
        ["Ngân sách cả nhóm (VNĐ)", budget],
        ["Dự toán cả nhóm (VNĐ)", total],
        ["Bình quân/người (VNĐ)", plan["per_person_vnd"]],
        ["Vượt ngân sách (VNĐ)", over_budget],
        ["Còn lại trong ngân sách (VNĐ)", max(0, budget - total)],
        ["Trạng thái ngân sách", "Vượt ngân sách" if over_budget else "Trong ngân sách"],
        ["Điều chỉnh bởi người dùng", "Có" if plan.get("manually_adjusted") else "Chưa"],
    ]
    _sheet(workbook, "Tổng quan", ["Thông tin", "Nội dung"], summary_rows, [36, 95])
    for row_number, (label, value) in enumerate(summary_rows, start=2):
        if label.endswith("(VNĐ)"):
            workbook["Tổng quan"].cell(row_number, 2).number_format = CURRENCY_FORMAT
    if over_budget:
        for row_number, (label, value) in enumerate(summary_rows, start=2):
            if label in ("Vượt ngân sách (VNĐ)", "Trạng thái ngân sách"):
                for cell in workbook["Tổng quan"][row_number]:
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")

    itinerary_rows, daily_cost_rows = [], []
    for day in plan["days"]:
        location, food, transport, costs = day["location"], day["food"], day["transport"], day["costs"]
        itinerary_rows.append([
            day["day"], location["name"], location["area"],
            f"Trung tâm Buôn Ma Thuột → {location['name']} → Trung tâm Buôn Ma Thuột",
            food["item"], transport["mode"], day["reason"], costs["total"], day["sample_id"],
        ])
        daily_cost_rows.append([
            day["day"], location["entrance_cost_vnd_per_person"], costs["tickets"],
            food["cost_vnd_per_person_per_meal"], request["meals_per_day"], costs["food"],
            transport["reference_party_size"], costs["transport_packages"],
            transport["cost_vnd_per_group_round_trip"], costs["transport"], costs["total"],
        ])
    _sheet(workbook, "Lịch trình", ["Ngày", "Điểm tham quan chính", "Khu vực", "Tuyến định hướng",
           "Món ăn tham khảo", "Phương tiện mẫu", "Lý do lựa chọn", "Tổng ngày, chưa gồm lưu trú", "Mã mẫu"],
           itinerary_rows, [9, 34, 22, 58, 34, 22, 65, 26, 14], (8,))
    _sheet(workbook, "Chi phí từng ngày", ["Ngày", "Vé/người", "Vé cả nhóm", "Đơn giá/bữa/người",
           "Số bữa/ngày", "Ăn uống cả nhóm", "Người/gói mẫu", "Số gói di chuyển",
           "Đơn giá/gói khứ hồi", "Di chuyển cả nhóm", "Tổng ngày, chưa gồm lưu trú"],
           daily_cost_rows, [9, 20, 22, 24, 16, 24, 19, 20, 24, 24, 28], (2, 3, 4, 6, 9, 10, 11))

    lodging = plan["lodging"]
    lodging_rows = ([[lodging["name"], lodging["area"], lodging["room_capacity_persons"], lodging["rooms"],
                     lodging["nights"], lodging["cost_vnd_per_room_per_night"], lodging["total"], lodging["sample_id"]]]
                    if lodging else [["Không tính lưu trú", "", 0, 0, 0, 0, 0, ""]])
    _sheet(workbook, "Lưu trú", ["Tên lưu trú giả lập", "Khu vực", "Người/phòng", "Số phòng", "Số đêm",
           "Đơn giá/phòng/đêm", "Tổng lưu trú cả chuyến", "Mã mẫu"],
           lodging_rows, [40, 24, 17, 14, 14, 26, 28, 14], (6, 7))
    cost_rows = [[label, plan["costs"][key]] for key, label in
                 [("tickets", "Vé tham quan"), ("food", "Ăn uống"), ("transport", "Di chuyển"), ("lodging", "Lưu trú")]]
    cost_rows.append(["Tổng dự toán cả nhóm", total])
    _sheet(workbook, "Tổng hợp chi phí", ["Hạng mục", "Số tiền"], cost_rows, [32, 28], (2,))
    notes = list(plan["warnings"])
    if over_budget:
        amount = f"{over_budget:,}".replace(",", ".")
        budget_note = f"Dự toán vượt ngân sách cả nhóm {amount} VNĐ."
        if not any("vượt ngân sách" in note.lower() for note in notes):
            notes.append(budget_note)
    notes.append("Lưu trú được tính một lần cho cả chuyến, không nằm trong tổng chi phí từng ngày.")
    notes.append("Bình quân/người được làm tròn lên đến 1 VNĐ.")
    _sheet(workbook, "Lưu ý", ["STT", "Nội dung"], [[index, note] for index, note in enumerate(notes, start=1)], [9, 110])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
